#!/usr/bin/env python3
"""One-command HTML deck export: QA -> 1920x1080 frames -> PDF + PPTX.

The accepted static final frame is the source of truth. The same captured PNGs
feed PDF and PPTX, so the three deliverables cannot drift visually.

Usage:
    python tools/export_deck.py deck.html
    python tools/export_deck.py deck.html --out dist/deck
    python tools/export_deck.py deck.html --skip-lint   # exceptional only
"""
from __future__ import annotations

import argparse
import base64
import hashlib
import json
import mimetypes
import os
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import quote, unquote, urlparse

W, H = 1920, 1080
W_IN, H_IN = 13.3333333333, 7.5

CAPTURE_CSS = r"""
html.__export,html.__export body{width:1920px!important;height:1080px!important;
min-width:1920px!important;min-height:1080px!important;margin:0!important;
padding:0!important;overflow:hidden!important;background:#000!important}
html.__export [data-stage]{position:absolute!important;inset:0!important;
width:1920px!important;height:1080px!important;transform:none!important;
transform-origin:top left!important;overflow:hidden!important}
html.__export [data-slide]{position:absolute!important;inset:0!important;
width:1920px!important;height:1080px!important;margin:0!important;opacity:0!important;
visibility:hidden!important;pointer-events:none!important;transition:none!important;
animation:none!important}
html.__export [data-slide].__export_target{opacity:1!important;visibility:visible!important}
html.__export .deck-shell,html.__export .deck-shell-progress,
html.__export .deck-shell-overview,html.__export .deck-shell-sr{display:none!important}
html.__export *,html.__export *::before,html.__export *::after{
transition:none!important;animation-play-state:paused!important}
"""

PREPARE_JS = r"""() => {
  document.documentElement.classList.add('__export','deck-capture');
  document.documentElement.classList.remove('deck-live','deck-ready');
  document.body.classList.remove('deck-live','deck-ready');
  document.querySelectorAll('*').forEach(el => [...el.classList].forEach(c => {
    if (c.startsWith('deck-step-') || /^phase-\d+$/.test(c) ||
        c === 'is-leaving' || c === 'is-cutting') el.classList.remove(c);
  }));
}"""

SELECT_JS = r"""i => {
  const slides=[...document.querySelectorAll('[data-slide]')];
  slides.forEach((s,n)=>s.classList.toggle('__export_target',n===i));
  const s=slides[i]; if(!s) return null;
  return {title:s.getAttribute('data-title') ||
    s.querySelector('h1,h2,[role="heading"]')?.textContent?.trim() || `Slide ${i+1}`,
    id:s.getAttribute('data-slide') || ''};
}"""


class ExportError(RuntimeError):
    pass


def repo_root() -> Path:
    return Path(__file__).resolve().parent.parent


def launch_kwargs() -> dict[str, str]:
    for env in ('PLAYWRIGHT_CHROMIUM_EXECUTABLE_PATH', 'CHROME_PATH'):
        p = os.environ.get(env)
        if p and Path(p).exists():
            return {'executable_path': p}
    for name in ('chromium', 'chromium-browser', 'google-chrome',
                 'google-chrome-stable', 'chrome'):
        p = shutil.which(name)
        if p:
            return {'executable_path': p}
    return {}


def deps():
    try:
        from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout
    except ImportError as e:
        raise ExportError('install: pip install playwright && playwright install chromium') from e
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as e:
        raise ExportError('install: pip install python-pptx') from e
    return sync_playwright, PWTimeout, Presentation, Inches


def install_local_route(page: Any, root: Path) -> None:
    root = root.resolve()
    def handle(route: Any, request: Any) -> None:
        u = urlparse(request.url)
        if u.hostname != 'deck.local':
            route.continue_(); return
        target = (root / unquote(u.path).lstrip('/')).resolve()
        try:
            target.relative_to(root)
        except ValueError:
            route.fulfill(status=403, body='forbidden'); return
        if not target.is_file():
            route.fulfill(status=404, body='not found'); return
        route.fulfill(status=200, path=str(target),
                      content_type=mimetypes.guess_type(target.name)[0] or 'application/octet-stream')
    page.route('http://deck.local/**', handle)


def html_with_base(source: Path, served_root: Path) -> str:
    if served_root == repo_root().resolve():
        parent = source.resolve().relative_to(served_root).parent.parts
        base = 'http://deck.local/' + ('/'.join(quote(x) for x in parent) + '/' if parent else '')
    else:
        base = 'http://deck.local/'
    raw = source.read_text(encoding='utf-8')
    tag = f'<base href="{base}">'
    pos = raw.lower().find('<head')
    if pos < 0:
        return tag + raw
    end = raw.find('>', pos)
    return raw[:end + 1] + tag + raw[end + 1:]


def run_lint(source: Path, out: Path) -> dict[str, Any]:
    lint = repo_root() / 'tools' / 'slide_lint.py'
    report, shots = out / 'lint.json', out / 'qa-shots'
    cmd = [sys.executable, str(lint), str(source), '--shots', str(shots), '--json', str(report)]
    print('[1/4] QA:', ' '.join(cmd))
    p = subprocess.run(cmd)
    if p.returncode >= 2:
        raise ExportError('slide QA failed with errors; export aborted')
    if not report.exists():
        raise ExportError('slide QA did not produce lint.json')
    json.loads(report.read_text(encoding='utf-8'))
    return {'returncode': p.returncode, 'report': report.name}


def png_ok(path: Path) -> None:
    b = path.read_bytes()[:24]
    if len(b) < 24 or b[:8] != b'\x89PNG\r\n\x1a\n':
        raise ExportError(f'invalid PNG: {path}')
    size = (int.from_bytes(b[16:20], 'big'), int.from_bytes(b[20:24], 'big'))
    if size != (W, H):
        raise ExportError(f'{path.name}: {size[0]}x{size[1]}, expected {W}x{H}')


def capture(source: Path, frames: Path, sync_playwright: Any, timeout: Any) -> list[dict[str, Any]]:
    frames.mkdir(parents=True, exist_ok=True)
    for p in frames.glob('slide-*.png'):
        p.unlink()
    root = repo_root().resolve()
    try:
        source.resolve().relative_to(root)
    except ValueError:
        root = source.resolve().parent
    print('[2/4] Capture static final frames')
    slides = []
    with sync_playwright() as pw:
        browser = pw.chromium.launch(**launch_kwargs())
        try:
            ctx = browser.new_context(viewport={'width': W, 'height': H},
                                      device_scale_factor=1, reduced_motion='reduce')
            page = ctx.new_page(); install_local_route(page, root)
            page.set_content(html_with_base(source, root), wait_until='load')
            try: page.evaluate('() => document.fonts ? document.fonts.ready : Promise.resolve()')
            except Exception: pass
            page.add_style_tag(content=CAPTURE_CSS); page.evaluate(PREPARE_JS); page.wait_for_timeout(120)
            n = page.evaluate("() => document.querySelectorAll('[data-slide]').length")
            if not n: raise ExportError('no [data-slide] elements found')
            for i in range(n):
                meta = page.evaluate(SELECT_JS, i)
                frame = frames / f'slide-{i+1:03d}.png'
                page.wait_for_timeout(40); page.screenshot(path=str(frame), animations='disabled')
                png_ok(frame)
                slides.append({'index': i + 1, 'title': meta['title'].strip(),
                               'data_slide': meta['id'], 'frame': frame.name})
            ctx.close()
        except timeout as e:
            raise ExportError(f'Chromium timed out while rendering {source.name}') from e
        finally:
            browser.close()
    return slides


def build_pdf(frames: Path, slides: list[dict[str, Any]], dest: Path, sync_playwright: Any) -> None:
    imgs = []
    for s in slides:
        data = base64.b64encode((frames / s['frame']).read_bytes()).decode('ascii')
        imgs.append(f'<section><img src="data:image/png;base64,{data}"></section>')
    html = f'''<!doctype html><style>
@page{{size:{W_IN}in {H_IN}in;margin:0}}*{{box-sizing:border-box}}html,body{{margin:0;padding:0}}
section{{width:{W_IN}in;height:{H_IN}in;break-after:page}}section:last-child{{break-after:auto}}
img{{display:block;width:100%;height:100%;object-fit:fill}}</style>{''.join(imgs)}'''
    print('[3/4] Build PDF')
    with sync_playwright() as pw:
        browser = pw.chromium.launch(**launch_kwargs())
        try:
            page = browser.new_page(); page.set_content(html, wait_until='load')
            page.pdf(path=str(dest), print_background=True, prefer_css_page_size=True,
                     margin={'top':'0','right':'0','bottom':'0','left':'0'})
        finally:
            browser.close()
    if not dest.exists() or dest.stat().st_size < 1024:
        raise ExportError(f'PDF export failed: {dest}')


def build_pptx(frames: Path, slides: list[dict[str, Any]], dest: Path,
               source: Path, Presentation: Any, Inches: Any) -> None:
    print('[4/4] Build PPTX')
    prs = Presentation(); prs.slide_width = Inches(W_IN); prs.slide_height = Inches(H_IN)
    blank = prs.slide_layouts[6]
    for s in slides:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(frames / s['frame']), 0, 0,
                                 width=prs.slide_width, height=prs.slide_height)
    prs.core_properties.title = source.stem
    prs.core_properties.subject = 'Static final-frame export from html-slide'
    prs.save(dest)
    if len(Presentation(dest).slides) != len(slides):
        raise ExportError('PPTX slide count mismatch')


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b''): h.update(chunk)
    return h.hexdigest()


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description='QA and export html-slide to PNG, PDF and PPTX')
    ap.add_argument('html', type=Path)
    ap.add_argument('--out', type=Path, default=None)
    ap.add_argument('--skip-lint', action='store_true',
                    help='exceptional only: skip tools/slide_lint.py')
    a = ap.parse_args(argv)
    source = a.html.resolve()
    out = (a.out or (repo_root() / 'dist' / a.html.stem)).resolve()
    try:
        if not source.is_file() or source.suffix.lower() not in ('.html', '.htm'):
            raise ExportError(f'invalid HTML deck: {source}')
        sync_playwright, timeout, Presentation, Inches = deps()
        out.mkdir(parents=True, exist_ok=True)
        lint = None if a.skip_lint else run_lint(source, out)
        if a.skip_lint: print('[1/4] QA: SKIPPED (--skip-lint)')
        frames = out / 'frames'
        slides = capture(source, frames, sync_playwright, timeout)
        pdf, pptx = out / f'{source.stem}.pdf', out / f'{source.stem}.pptx'
        build_pdf(frames, slides, pdf, sync_playwright)
        build_pptx(frames, slides, pptx, source, Presentation, Inches)
        manifest = {
            'source': str(source), 'source_sha256': sha256(source),
            'created_utc': datetime.now(timezone.utc).isoformat(),
            'stage': {'width': W, 'height': H, 'aspect': '16:9'},
            'slide_count': len(slides), 'slides': slides,
            'pdf': pdf.name, 'pptx': pptx.name, 'lint': lint,
        }
        (out / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
        print(f'\nEXPORT PASS\n  slides: {len(slides)}\n  pdf:    {pdf}\n  pptx:   {pptx}\n  frames: {frames}')
        return 0
    except ExportError as e:
        print(f'EXPORT FAIL: {e}', file=sys.stderr); return 2
    except Exception as e:
        print(f'EXPORT FAIL: unexpected error: {e}', file=sys.stderr); return 2


if __name__ == '__main__':
    raise SystemExit(main())
